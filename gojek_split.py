#!/usr/bin/env python3
"""
gojek_split.py - Pisahkan PDF "Riwayat transaksi Gojek" per layanan (GoFood, GoSend, dst).

Pemakaian:
    python gojek_split.py Riwayat_transaksi_Gojek.pdf
    python gojek_split.py Riwayat_transaksi_Gojek.pdf --only GoFood GoSend
    python gojek_split.py Riwayat_transaksi_Gojek.pdf -o hasil --excel

Instalasi:
    pip install pdfplumber reportlab openpyxl     (openpyxl hanya untuk --excel)

Catatan: parser membaca PDF unduhan aplikasi Gojek berdasarkan posisi kolom
(Tanggal | No. transaksi | Layanan | Dari | Tujuan | Metode bayar | Total dibayar).
Jika Gojek mengubah format PDF, parser perlu disesuaikan. Selalu cek peringatan
rekonsiliasi total yang dicetak di akhir proses.
"""
import argparse
import re
import sys
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

import pdfplumber
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

DATE_RE = re.compile(r"^\d{2}/\d{2}/\d{4}$")
HEADERS = ["Tanggal", "No.", "Layanan", "Dari", "Tujuan", "Metode", "Total"]


def rp(n: int) -> str:
    return "Rp" + f"{n:,}".replace(",", ".")


def parse_rp(text: str):
    m = re.search(r"Rp\s*([\d.]+)", text)
    return int(m.group(1).replace(".", "")) if m else None


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #
def parse_pdf(path: str):
    """Return (info, records). Setiap record = dict transaksi."""
    with pdfplumber.open(path) as pdf:
        page1 = pdf.pages[0]
        w1 = page1.extract_words()

        # --- posisi kolom dari baris header tabel ---
        lay = next((w for w in w1 if w["text"] == "Layanan"), None)
        if lay is None:
            raise ValueError("Header tabel 'Layanan' tidak ditemukan. Apakah ini PDF riwayat transaksi Gojek?")
        hdr_words = [w for w in w1 if abs(w["top"] - lay["top"]) < 8 and w["text"] in HEADERS]
        col_x = OrderedDict()
        for h in HEADERS:
            w = next((w for w in hdr_words if w["text"] == h), None)
            if w is None:
                raise ValueError(f"Kolom '{h}' tidak ditemukan di header.")
            col_x[h] = w["x0"]
        bounds = [x - 2 for x in col_x.values()]  # batas kiri tiap kolom
        table_top = lay["top"] + 16  # lewati baris header 2 baris

        # --- info akun & periode & total ---
        text1 = page1.extract_text() or ""
        periode = re.search(r"(\d{1,2}\s+\w+\s+\d{4})\s*-\s*(\d{1,2}\s+\w+\s+\d{4})", text1)
        total_hdr = next((w for w in w1 if re.match(r"^Rp[\d.]+$", w["text"]) and w["top"] < lay["top"]), None)
        info = {
            "periode": f"{periode.group(1)} - {periode.group(2)}" if periode else "",
            "total_header": parse_rp(total_hdr["text"]) if total_hdr else None,
            "nama": "", "telp": "", "email": "",
        }
        left = [w for w in w1 if w["x0"] < 200 and w["top"] < lay["top"] - 20 and w["top"] > 60]
        lines = {}
        for w in left:
            lines.setdefault(round(w["top"] / 3), []).append(w["text"])
        lines = [" ".join(v) for _, v in sorted(lines.items())]
        for ln in lines:
            if "@" in ln:
                info["email"] = ln
            elif re.match(r"^\+?\d{8,}$", ln.replace(" ", "")):
                info["telp"] = ln
            elif not info["nama"]:
                info["nama"] = ln

        # --- kumpulkan kata tabel per halaman ---
        def col_of(x0):
            idx = 0
            for i, b in enumerate(bounds):
                if x0 >= b:
                    idx = i
            return idx

        records, cur = [], None
        for pno, page in enumerate(pdf.pages):
            words = page.extract_words(extra_attrs=["fontname"])
            words.sort(key=lambda w: (round(w["top"] / 2), w["x0"]))
            for w in words:
                if pno == 0 and w["top"] < table_top:
                    continue
                c = col_of(w["x0"])
                if c == 0 and DATE_RE.match(w["text"]):
                    cur = {i: [] for i in range(len(HEADERS))}
                    records.append(cur)
                if cur is None:
                    continue
                cur[c].append((w["text"], "Bold" in w["fontname"]))

    return info, [_build_record(r) for r in records]


def _segments(words):
    """Pecah kata menjadi (is_bold, teks) berurutan."""
    segs = []
    for text, bold in words:
        if segs and segs[-1][0] == bold:
            segs[-1][1].append(text)
        else:
            segs.append((bold, [text]))
    return [(b, " ".join(t)) for b, t in segs]


def _name_addr(words):
    segs = _segments(words)
    name = " ".join(t for b, t in segs if b)
    addr = " ".join(t for b, t in segs if not b)
    if not name:  # jaga-jaga bila tidak ada teks bold
        name, addr = addr, ""
    return name, addr


def _build_record(cells):
    tgl_words = [t for t, _ in cells[0]]
    tanggal = tgl_words[0]
    jam = " ".join(tgl_words[1:])
    no = "".join(t for t, _ in cells[1])
    lay_words = [t for t, _ in cells[2]]
    layanan = lay_words[0] if lay_words else "?"
    tier = " ".join(lay_words[1:])
    dari = _name_addr(cells[3])
    tujuan = _name_addr(cells[4])
    metode = " ".join(t for t, _ in cells[5])
    total = parse_rp(" ".join(t for t, _ in cells[6]))
    return {"tanggal": tanggal, "jam": jam, "no": no, "layanan": layanan, "tier": tier,
            "dari": dari, "tujuan": tujuan, "metode": metode, "total": total}


# --------------------------------------------------------------------------- #
# Output PDF (landscape, gaya laporan Gojek yang dirapikan)
# --------------------------------------------------------------------------- #
GREEN = colors.HexColor("#00880F")
GREEN_DARK = colors.HexColor("#0B5D1E")
LIGHT = colors.HexColor("#E8F7EC")
ZEBRA = colors.HexColor("#F7FBF8")
LINE = colors.HexColor("#DADADA")
INK = colors.HexColor("#1F1F1F")
MUTED = colors.HexColor("#666666")
MARGIN = 12 * mm

# True = alamat diringkas di PDF (buang awalan yang sama dengan nama tempat dan ", Indonesia").
# File Excel selalu memuat alamat lengkap apa adanya.
CLEAN_ADDRESS = True


def _clean_addr(name: str, addr: str) -> str:
    a = addr or ""
    if CLEAN_ADDRESS:
        if name and a.lower().startswith(name.lower()):
            a = a[len(name):].lstrip(" ,")
        a = re.sub(r",?\s*Indonesia\.?$", "", a).strip(" ,")
    return a


class _NumberedCanvas(rl_canvas.Canvas):
    """Canvas yang bisa mencetak 'Halaman X dari Y' dan footer."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved = []

    def showPage(self):
        self._saved.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._saved)
        for state in self._saved:
            self.__dict__.update(state)
            self._footer(total)
            super().showPage()
        super().save()

    def _footer(self, total):
        w, _ = self._pagesize
        self.setStrokeColor(LINE)
        self.setLineWidth(0.5)
        self.line(MARGIN, 11 * mm, w - MARGIN, 11 * mm)
        self.setFont("Helvetica", 7)
        self.setFillColor(MUTED)
        self.drawString(MARGIN, 7 * mm,
                        f"Dibuat otomatis dari riwayat transaksi Gojek - {datetime.now():%d/%m/%Y %H:%M}")
        self.drawRightString(w - MARGIN, 7 * mm, f"Halaman {self._pageNumber} dari {total}")


def write_pdf(path, layanan, rows, info, note=None):
    page = landscape(A4)
    total = sum(r["total"] or 0 for r in rows)
    avg = round(total / len(rows)) if rows else 0
    e = lambda s: escape(s or "")

    def header(c, d):
        w, h = page
        c.saveState()
        c.setFillColor(INK)
        c.rect(0, h - 13 * mm, w, 13 * mm, fill=1, stroke=0)
        c.setFillColor(GREEN)
        c.rect(0, h - 13.8 * mm, w, 0.8 * mm, fill=1, stroke=0)
        c.circle(MARGIN + 2.6 * mm, h - 6.5 * mm, 2.6 * mm, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.circle(MARGIN + 2.6 * mm, h - 6.5 * mm, 1.1 * mm, fill=1, stroke=0)
        c.setFont("Helvetica-Bold", 13)
        c.drawString(MARGIN + 7.5 * mm, h - 8.3 * mm, "gojek")
        c.setFont("Helvetica", 9)
        c.drawRightString(w - MARGIN, h - 8 * mm, f"Riwayat Transaksi {layanan}")
        c.restoreState()

    doc = SimpleDocTemplate(path if hasattr(path, "write") else str(path), pagesize=page, leftMargin=MARGIN, rightMargin=MARGIN,
                            topMargin=21 * mm, bottomMargin=16 * mm,
                            title=f"Riwayat Transaksi Gojek - {layanan}", author=info.get("nama") or "")

    st = lambda name, **kw: ParagraphStyle(name, fontName=kw.pop("fontName", "Helvetica"), **kw)
    body = st("body", fontSize=8, leading=10, textColor=INK)
    bold = st("bold", fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=INK)
    addr = st("addr", fontSize=7, leading=9, textColor=MUTED)
    th = st("th", fontName="Helvetica-Bold", fontSize=8, leading=10, textColor=colors.white)
    th_r = ParagraphStyle("thr", parent=th, alignment=2)
    amt = ParagraphStyle("amt", parent=bold, alignment=2)
    lbl = st("lbl", fontSize=7, leading=9, textColor=MUTED)
    val = st("val", fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=INK)
    val_s = st("vals", fontName="Helvetica-Bold", fontSize=9.5, leading=15, textColor=INK)
    val_g = st("valg", fontName="Helvetica-Bold", fontSize=13, leading=15, textColor=GREEN)
    acc_n = st("accn", fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=INK)

    # ---- kartu ringkasan ----
    account = [Paragraph(e(info.get("nama")), acc_n), Paragraph(e(info.get("telp")), body),
               Paragraph(e(info.get("email")), body)]
    cards = [account,
             [Paragraph("Periode transaksi", lbl), Paragraph(e(info.get("periode")), val_s)],
             [Paragraph("Jumlah transaksi", lbl), Paragraph(str(len(rows)), val)],
             [Paragraph(f"Total {e(layanan)}", lbl), Paragraph(rp(total), val_g)],
             [Paragraph("Rata-rata per transaksi", lbl), Paragraph(rp(avg), val)]]
    widths = [64 * mm, 64 * mm, 40 * mm, 47 * mm, 46 * mm]
    gap = 3 * mm
    row, cw = [], []
    for i, (c, w) in enumerate(zip(cards, widths)):
        if i:
            row.append("")
            cw.append(gap)
        row.append(c)
        cw.append(w)
    summary = Table([row], colWidths=cw)
    cmds = [("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7), ("LEFTPADDING", (0, 0), (-1, -1), 9)]
    for col in range(0, len(row), 2):
        if col == 0:
            cmds.append(("BOX", (col, 0), (col, 0), 0.6, LINE))
        else:
            cmds += [("BACKGROUND", (col, 0), (col, 0), LIGHT), ("BOX", (col, 0), (col, 0), 0.6, LIGHT)]
    cmds += [("LEFTPADDING", (1, 0), (1, 0), 0), ("RIGHTPADDING", (1, 0), (1, 0), 0)]
    summary.setStyle(TableStyle(cmds))

    story = [summary, Spacer(1, 10),
             Paragraph(f"Rincian transaksi <b>{e(layanan)}</b>. Bukti satuan dari semua transaksi dapat "
                       f"diunduh di aplikasi Gojek.", lbl)]
    if note:  # mis. penanda bahwa data sudah disesuaikan manual
        story.append(Paragraph(f"<b>Catatan:</b> {e(note)}",
                               st("note", fontSize=7.5, leading=10, textColor=colors.HexColor("#B45309"))))
    story.append(Spacer(1, 6))

    # ---- tabel transaksi ----
    data = [[Paragraph("Tanggal", th), Paragraph("No. transaksi", th), Paragraph("Layanan", th),
             Paragraph("Dari", th), Paragraph("Tujuan", th), Paragraph("Metode bayar", th),
             Paragraph("Total dibayar", th_r)]]
    for r in rows:
        lay = f"<b>{e(r['layanan'])}</b>" + (f"<br/><font size='7' color='#666666'>{e(r['tier'])}</font>"
                                             if r["tier"] else "")

        def place(p):
            cell = [Paragraph(e(p[0]), bold)]
            a = _clean_addr(p[0], p[1])
            if a:
                cell += [Spacer(1, 1.5), Paragraph(e(a), addr)]
            return cell

        data.append([
            Paragraph(f"<b>{e(r['tanggal'])}</b><br/><font size='7' color='#666666'>{e(r['jam'])}</font>", body),
            Paragraph(e(r["no"]), body), Paragraph(lay, body),
            place(r["dari"]), place(r["tujuan"]),
            Paragraph(e(r["metode"]), body), Paragraph(rp(r["total"] or 0), amt)])
    data.append(["", "", "", "", "", Paragraph("Total", bold), Paragraph(rp(total), amt)])

    t = Table(data, colWidths=[26 * mm, 32 * mm, 22 * mm, 68 * mm, 68 * mm, 22 * mm, 35 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), GREEN_DARK),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, ZEBRA]),
        ("LINEBELOW", (0, 1), (-1, -2), 0.4, LINE),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("BACKGROUND", (0, -1), (-1, -1), LIGHT), ("LINEABOVE", (0, -1), (-1, -1), 1, GREEN),
        ("BOX", (0, 0), (-1, -1), 0.5, LINE)]))
    story.append(t)
    doc.build(story, onFirstPage=header, onLaterPages=header, canvasmaker=_NumberedCanvas)
    return total


def write_excel(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.append(["Tanggal", "Jam", "No. Transaksi", "Layanan", "Dari", "Alamat Dari", "Tujuan",
               "Alamat Tujuan", "Metode Bayar", "Total Dibayar"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8F7EC")
    for r in rows:
        ws.append([r["tanggal"], r["jam"], r["no"], r["layanan"], r["dari"][0], r["dari"][1],
                   r["tujuan"][0], r["tujuan"][1], r["metode"], r["total"]])
    n = len(rows) + 1
    ws.append([None] * 8 + ["TOTAL", f"=SUM(J2:J{n})"])
    ws.cell(row=n + 1, column=9).font = Font(bold=True)
    ws.cell(row=n + 1, column=10).font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=10, max_col=10):
        row[0].number_format = '"Rp"#,##0'
    for col, wdt in zip("ABCDEFGHIJ", [12, 12, 18, 10, 30, 50, 30, 50, 14, 16]):
        ws.column_dimensions[col].width = wdt
    wb.save(path)


def pdf_bytes(layanan, rows, info, note=None) -> bytes:
    """PDF laporan sebagai bytes (untuk aplikasi web / unduhan)."""
    import io
    buf = io.BytesIO()
    write_pdf(buf, layanan, rows, info, note)
    return buf.getvalue()


def excel_bytes(rows) -> bytes:
    """Excel laporan sebagai bytes (untuk aplikasi web / unduhan)."""
    import io
    buf = io.BytesIO()
    write_excel(buf, rows)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Main
# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description="Pisahkan riwayat transaksi Gojek (PDF) per layanan.")
    ap.add_argument("pdf", help="File PDF riwayat transaksi Gojek")
    ap.add_argument("-o", "--out", default=None, help="Folder output (default: folder file input)")
    ap.add_argument("--only", nargs="+", metavar="LAYANAN",
                    help="Hanya layanan tertentu, mis. --only GoFood GoSend (default: semua layanan)")
    ap.add_argument("--excel", action="store_true", help="Buat juga file Excel per layanan")
    args = ap.parse_args()

    src = Path(args.pdf)
    out = Path(args.out) if args.out else src.parent
    out.mkdir(parents=True, exist_ok=True)

    info, records = parse_pdf(str(src))
    if not records:
        sys.exit("Tidak ada transaksi terbaca dari PDF.")

    groups = OrderedDict()
    for r in records:
        groups.setdefault(r["layanan"], []).append(r)

    only = {s.lower() for s in args.only} if args.only else None
    suffix = re.sub(r"^Riwayat_transaksi_Gojek_?", "", src.stem)

    print(f"Terbaca {len(records)} transaksi, periode {info['periode']}")
    grand = 0
    for lay, rows in groups.items():
        sub = sum(r["total"] or 0 for r in rows)
        grand += sub
        tag = "" if (only is None or lay.lower() in only) else "  (dilewati)"
        print(f"  {lay:<8} {len(rows):>3} transaksi  {rp(sub):>14}{tag}")
        if only is not None and lay.lower() not in only:
            continue
        name = f"Riwayat_transaksi_Gojek_{lay}" + (f"_{suffix}" if suffix else "")
        write_pdf(out / f"{name}.pdf", lay, rows, info)
        if args.excel:
            write_excel(out / f"{name}.xlsx", rows)

    # rekonsiliasi dengan total di header PDF
    missing = [r["no"] for r in records if r["total"] is None]
    if missing:
        print(f"PERINGATAN: {len(missing)} transaksi tanpa nominal terbaca: {missing}")
    if info["total_header"] is not None:
        if grand == info["total_header"]:
            print(f"OK: total semua transaksi {rp(grand)} sama dengan total di header PDF.")
        else:
            print(f"PERINGATAN: total terbaca {rp(grand)} != total header {rp(info['total_header'])}. "
                  f"Periksa hasil parsing sebelum dipakai.")
    print(f"File tersimpan di: {out.resolve()}")


if __name__ == "__main__":
    main()